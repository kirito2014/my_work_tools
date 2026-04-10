import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import threading
import logging
from openpyxl import load_workbook

# ==========================================
# 资源路径处理函数 (兼容本地运行与打包后的EXE)
# ==========================================
def get_resource_path(relative_path):
    """获取资源文件的绝对路径"""
    if hasattr(sys, '_MEIPASS'):
        # PyInstaller 打包后的临时目录
        return os.path.join(sys._MEIPASS, relative_path)
    # 本地开发环境目录
    return os.path.join(os.path.abspath("."), relative_path)


# 导入我们刚刚写的核心数据处理类
try:
    from gather_case_list import DataAnalyzer
except ImportError:
    messagebox.showerror("错误", "找不到 gather_case_list.py 文件，请确保它们在同一目录下！")
    sys.exit(1)

# ==========================================
# 自定义日志处理器：将日志发送到 Tkinter Text 组件
# ==========================================
class TextHandler(logging.Handler):
    def __init__(self, text_widget):
        logging.Handler.__init__(self)
        self.text_widget = text_widget
        
        # 配置关键字的彩色 Tag
        self.text_widget.tag_config("INFO", foreground="black")
        self.text_widget.tag_config("WARNING", foreground="#D2691E") # 巧克力色
        self.text_widget.tag_config("ERROR", foreground="red")
        self.text_widget.tag_config("SUCCESS", foreground="green")
        self.text_widget.tag_config("HIGHLIGHT", foreground="blue", font=("Consolas", 10, "bold"))

    def emit(self, record):
        msg = self.format(record)
        
        tag = "INFO"
        if record.levelno == logging.WARNING:
            tag = "WARNING"
        elif record.levelno >= logging.ERROR:
            tag = "ERROR"
            
        if "成功" in msg or "完美" in msg or "✅" in msg:
            tag = "SUCCESS"

        def append():
            self.text_widget.configure(state='normal')
            self.text_widget.insert(tk.END, msg + '\n', tag)
            self.text_widget.configure(state='disabled')
            self.text_widget.yview(tk.END) 
            
        self.text_widget.after(0, append)


class DataAnalyzerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("数据统计分析自动化工具 v1.0")
        self.root.geometry("750x700")
        self.root.resizable(False, False)
        
        # 加载窗口图标
        icon_path = get_resource_path(os.path.join("icons", "sunline.ico"))
        if os.path.exists(icon_path):
            self.root.iconbitmap(icon_path)
        else:
            print(f"Warning: 找不到图标文件 {icon_path}")
        
        self.source_dir = tk.StringVar()
        self.target_file = tk.StringVar()
        self.is_running = False

        self.setup_ui()
        self.setup_logging()

    def setup_ui(self):
        frame_top = tk.Frame(self.root, padx=15, pady=15)
        frame_top.pack(fill=tk.X)

        tk.Button(frame_top, text="选择源文件夹", command=self.select_source, width=15).grid(row=0, column=0, pady=5, sticky="w")
        tk.Entry(frame_top, textvariable=self.source_dir, width=50, state='readonly').grid(row=0, column=1, padx=10, pady=5)
        self.lbl_source_cnt = tk.Label(frame_top, text="文件数量: 0", fg="blue", font=("微软雅黑", 10, "bold"))
        self.lbl_source_cnt.grid(row=0, column=2, pady=5)

        tk.Button(frame_top, text="选择目标文件", command=self.select_target, width=15).grid(row=1, column=0, pady=5, sticky="w")
        tk.Entry(frame_top, textvariable=self.target_file, width=50, state='readonly').grid(row=1, column=1, padx=10, pady=5)
        self.lbl_target_cnt = tk.Label(frame_top, text="表数量: 0", fg="blue", font=("微软雅黑", 10, "bold"))
        self.lbl_target_cnt.grid(row=1, column=2, pady=5)

        frame_mid = tk.LabelFrame(self.root, text="执行日志", padx=10, pady=10)
        frame_mid.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)

        self.log_area = scrolledtext.ScrolledText(frame_mid, wrap=tk.WORD, font=("Consolas", 10), state='disabled', bg="#F8F8F8")
        self.log_area.pack(fill=tk.BOTH, expand=True)

        frame_bottom = tk.Frame(self.root, padx=15, pady=10)
        frame_bottom.pack(fill=tk.X)

        self.btn_start = tk.Button(frame_bottom, text="▶ 开始执行", command=self.start_execution, bg="#4CAF50", fg="white", font=("微软雅黑", 10, "bold"), width=15)
        self.btn_start.pack(side=tk.LEFT, padx=5)

        self.btn_exit = tk.Button(frame_bottom, text="✖ 退出程序", command=self.root.quit, bg="#f44336", fg="white", font=("微软雅黑", 10, "bold"), width=15)
        self.btn_exit.pack(side=tk.LEFT, padx=5)

    def setup_logging(self):
        self.app_logger = logging.getLogger("gather_case_list")
        self.app_logger.setLevel(logging.INFO)
        text_handler = TextHandler(self.log_area)
        formatter = logging.Formatter('%(asctime)s - %(message)s', datefmt='%H:%M:%S')
        text_handler.setFormatter(formatter)
        self.app_logger.addHandler(text_handler)

    def select_source(self):
        folder = filedialog.askdirectory(title="选择包含源XLSX的文件夹")
        if folder:
            self.source_dir.set(folder)
            try:
                count = len([f for f in os.listdir(folder) if f.endswith('.xlsx') and not f.startswith('~')])
                self.lbl_source_cnt.config(text=f"文件数量: {count}")
            except Exception as e:
                self.lbl_source_cnt.config(text="读取错误", fg="red")

    def select_target(self):
        file_path = filedialog.askopenfilename(title="选择目标结果文件", filetypes=[("Excel Files", "*.xlsx")])
        if file_path:
            self.target_file.set(file_path)
            self.lbl_target_cnt.config(text="正在读取...", fg="orange")
            self.root.update()
            threading.Thread(target=self._count_target_rows, args=(file_path,), daemon=True).start()

    def _count_target_rows(self, file_path):
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            if "SIT2阶段-明细进度" not in wb.sheetnames:
                self.lbl_target_cnt.config(text="缺失指定Sheet", fg="red")
                return
                
            ws = wb["SIT2阶段-明细进度"]
            table_count = 0
            
            for row_idx in range(3, ws.max_row + 1):
                cell_val = ws.cell(row=row_idx, column=4).value
                if cell_val is not None and str(cell_val).strip() != "":
                    table_count += 1
                    
            self.lbl_target_cnt.config(text=f"表数量: {table_count}", fg="blue")
            wb.close()
        except Exception as e:
            self.lbl_target_cnt.config(text="读取失败", fg="red")

    def log_msg(self, msg, tag="INFO"):
        self.log_area.configure(state='normal')
        self.log_area.insert(tk.END, msg + '\n', tag)
        self.log_area.configure(state='disabled')
        self.log_area.yview(tk.END)

    def start_execution(self):
        if self.is_running:
            return
            
        src = self.source_dir.get()
        tgt = self.target_file.get()
        
        if not src or not tgt:
            messagebox.showwarning("提示", "请先选择源文件夹和目标文件！")
            return
            
        self.is_running = True
        self.btn_start.config(text="⏳ 执行中...", state=tk.DISABLED, bg="gray")
        self.log_area.configure(state='normal')
        self.log_area.delete(1.0, tk.END)
        self.log_area.configure(state='disabled')
        
        self.log_msg("=== 启动自动化处理任务 ===", "HIGHLIGHT")
        
        threading.Thread(target=self._run_analyzer, args=(src, tgt), daemon=True).start()

    def _run_analyzer(self, src, tgt):
        try:
            analyzer = DataAnalyzer(
                source_dir=src,
                target_file=tgt,
                config_file='config.ini' 
            )
            analyzer.load_config()
            analyzer.load_source_files()
            analyzer.process_data()
            
            self.log_msg("\n=== 任务全部处理完毕！===", "SUCCESS")
            messagebox.showinfo("完成", "数据统计及格式渲染已全部完成！")
            
        except Exception as e:
            self.log_msg(f"\n[致命错误] 程序中断: {str(e)}", "ERROR")
            messagebox.showerror("错误", f"执行过程中发生错误：\n{str(e)}")
            
        finally:
            self.is_running = False
            self.btn_start.config(text="▶ 开始执行", state=tk.NORMAL, bg="#4CAF50")

if __name__ == "__main__":
    root = tk.Tk()
    app = DataAnalyzerGUI(root)
    root.mainloop()