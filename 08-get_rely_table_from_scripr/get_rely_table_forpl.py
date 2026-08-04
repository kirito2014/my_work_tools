# -*- coding:utf-8 -*-
# 脚本功能：遍历指定文件夹（含子文件夹），分析所有 .pl 脚本，
# 提取脚本的英文表名/中文表名/开发人员/目标库，以及该脚本所依赖的来源库名和来源表名，
# 最终写出一份 Excel 清单。
import os
import sys
import re
import openpyxl

# ----------------------------------------------------------------------------------------------
# 库变量默认值映射：脚本正文中经常出现 ${RRSDB}、${GDMDB}、${ODSDB} 这类 Perl 变量作为 schema，
# 这里维护一份"变量名 -> 真实库名"的兜底映射，用于统一替换。
# 如果某个 .pl 脚本里自己定义了同名变量（如 my $RRSDB=$ENV{"AUTO_RRSDB"}||'mrrs';），
# 会优先使用脚本内解析出来的值；只有脚本内没有解析到时才使用下面的兜底值。
# 如后续发现新的库变量，直接在这里补充即可。
# ----------------------------------------------------------------------------------------------
DEFAULT_DB_VAR_MAP = {
    'ODSDB': 'odsdata',   # ODS层库名
    'GDMDB': 'gdmdata',   # GDM层库名
    'RRSDB': 'mrrs',      # RRS层库名（大多数脚本的目标表都在这个库）
}

# 只在脚本头部注释区域内查找 Target Table / Comment / Developed By 信息，避免误匹配到脚本正文
HEADER_FIELD_MAX_LINES = 30


def normalize_label(label):
    """把注释里的字段标签统一处理成 只保留大写字母，便于做模糊匹配（忽略空格、大小写差异）"""
    return re.sub(r'[^A-Z]', '', label.upper())


def extract_header_info(file_path):
    """
    从脚本头部注释中提取：
      - comment      : "Comment" 字段的值，作为中文表名（找不到则为空字符串）
      - developed_by : "Developed By" 字段的值，作为开发人员（找不到则默认"无"）
    """
    comment = ''
    developed_by = ''
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            for i, line in enumerate(f):
                if i >= HEADER_FIELD_MAX_LINES:
                    break
                if ':' not in line:
                    continue
                clean_line = line.strip().lstrip('#').strip().lstrip('/*').strip()
                if ':' not in clean_line:
                    continue
                label, _, value = clean_line.partition(':')
                label_norm = normalize_label(label)
                value = value.strip()
                value = re.sub(r'\*/\s*$', '', value).strip()  # 去掉行尾的 */ 注释结束符
                if label_norm == 'COMMENT' and not comment:
                    comment = value
                elif label_norm == 'DEVELOPEDBY' and not developed_by:
                    developed_by = value
    except Exception as e:
        print(f"处理{file_path}获取脚本头部信息失败: {e}")

    if not developed_by:
        developed_by = '无'
    return comment, developed_by


def extract_db_var_map(raw_content):
    """
    从脚本正文中动态解析库变量定义，例如：
      my $RRSDB=$ENV{"AUTO_RRSDB"}||'mrrs';
    返回 {变量名: 默认值}，用于后续把 SQL 里的 ${变量名} 替换成真实库名。

    注意：部分手写脚本存在笔误，把默认值写成了自引用占位符，例如：
      my $GDMDB=$ENV{"AUTO_GDMDB"}||'${GDMDB}';
    这种默认值本身是个变量引用而不是真实库名，直接使用会导致替换后仍然是 ${GDMDB}，
    无法被后续正则识别为真实的 schema。因此这里会丢弃这种自引用占位符，
    转而使用脚本开头 DEFAULT_DB_VAR_MAP 里的兜底真实库名。
    """
    var_map = {}
    pattern = r'\$(\w+)\s*=\s*\$ENV\{[^}]*\}\s*\|\|\s*[\'"]([^\'"]*)[\'"]'
    for var_name, default_val in re.findall(pattern, raw_content):
        if default_val.strip().startswith('$'):
            continue
        var_map[var_name] = default_val
    return var_map


def substitute_db_vars(raw_content, var_map):
    """把脚本正文中的 ${变量名} 统一替换为真实库名（脚本内解析到的优先，找不到则用默认映射）"""
    merged_map = dict(DEFAULT_DB_VAR_MAP)
    merged_map.update(var_map)

    content = raw_content.replace('${version_num}', '')  # 版本号占位符直接去掉
    for var_name, real_val in merged_map.items():
        content = content.replace('${%s}' % var_name, real_val)
    return content


def determine_target_info(content, file_name):
    """
    确定"英文表名"、"目标库"、"目标语句类型"：
      1. 优先在脚本正文（已完成库变量替换）中查找下面几类"真正落地"的语句
         （schema 前缀可选，即 schema.table 或只写 table 都能识别），排除 VT_ 开头的临时表：
           - INSERT INTO <table> / INSERT OVERWRITE TABLE <table>  → 目标语句类型记为 'INSERT'
           - CREATE VIEW <table> / CREATE OR REPLACE VIEW <table>  → 目标语句类型记为 'VIEW'
         取脚本中第一个命中的语句：
           - table 即为英文表名；
           - 如果语句里写了 schema，则该 schema 即为目标库；
           - 如果语句里没写 schema，目标库标记为"未匹配到目标库"，但英文表名依然能正确取到。
      2. 如果脚本里完全找不到这样的语句，
         则英文表名使用脚本文件名（去掉 .pl 后缀）作为备选，目标库标记为"未匹配到目标库"，
         目标语句类型记为 None，需要人工核查该脚本。
    """
    pattern = (r'(INSERT\s+(?:INTO|OVERWRITE\s+TABLE)|CREATE\s+(?:OR\s+REPLACE\s+)?VIEW)'
               r'\s+(?:(\w+)\.)?(\w+)')
    matches = re.findall(pattern, content, flags=re.IGNORECASE)

    for keyword, schema, table in matches:
        if table.upper().startswith('VT_'):
            continue
        schema_out = schema.upper() if schema else '未匹配到目标库'
        target_type = 'VIEW' if keyword.strip().upper().startswith('CREATE') else 'INSERT'
        return table.upper(), schema_out, target_type

    default_name = os.path.splitext(file_name)[0].upper()
    return default_name, '未匹配到目标库', None


def determine_process_type(content, target_type):
    """
    根据脚本内容判断对应目标表所属的"处理属性"，主要依据 determine_target_info 匹配到的语句类型，
    并结合脚本里出现的其它典型语句做进一步细化，方便快速了解每个脚本/表在做什么：
      - 视图处理           : 命中 CREATE VIEW / CREATE OR REPLACE VIEW
      - 实体加工(先删后插)  : 命中 INSERT，且脚本里也有 DELETE FROM（先清理再插入的典型写法）
      - 实体加工(先清空后插): 命中 INSERT，且脚本里有 TRUNCATE TABLE
      - 实体加工(直接插入/追加) : 命中 INSERT，但没有配套的 DELETE/TRUNCATE
      - 实体加工(仅更新)    : 没有 INSERT/CREATE VIEW，但有 UPDATE ... SET 语句
      - 实体加工(仅删除)    : 没有 INSERT/CREATE VIEW，但有 DELETE FROM 语句
      - 表结构维护(建表)    : 没有以上任何操作，但有 CREATE TABLE 语句
      - 表结构维护(变更)    : 没有以上任何操作，但有 ALTER TABLE 语句
      - 变量赋值/参数取值   : 只有 SELECT 语句，没有任何落地/变更类语句（如脚本里单纯查日期维表取变量）
      - 未识别             : 以上都没有命中，需要人工核查
    如果脚本里同时出现 GRANT 授权语句，会在分类结果后追加"+权限授予"标识。
    """
    has_delete = bool(re.search(r'DELETE\s+FROM\s+', content, flags=re.IGNORECASE))
    has_truncate = bool(re.search(r'TRUNCATE\s+TABLE\s+', content, flags=re.IGNORECASE))
    has_update = bool(re.search(r'UPDATE\s+\S+\s+SET\s+', content, flags=re.IGNORECASE))
    has_create_table = bool(re.search(r'CREATE\s+TABLE\s+', content, flags=re.IGNORECASE))
    has_alter_table = bool(re.search(r'ALTER\s+TABLE\s+', content, flags=re.IGNORECASE))
    has_select = bool(re.search(r'\bSELECT\b', content, flags=re.IGNORECASE))
    has_grant = bool(re.search(r'\bGRANT\b', content, flags=re.IGNORECASE))

    if target_type == 'VIEW':
        label = '视图处理'
    elif target_type == 'INSERT':
        if has_delete:
            label = '实体加工(先删后插)'
        elif has_truncate:
            label = '实体加工(先清空后插)'
        else:
            label = '实体加工(直接插入/追加)'
    else:
        # 没有匹配到 INSERT / CREATE VIEW 时，进一步根据脚本里出现的其它语句判断
        if has_update:
            label = '实体加工(仅更新)'
        elif has_delete:
            label = '实体加工(仅删除)'
        elif has_create_table:
            label = '表结构维护(建表)'
        elif has_alter_table:
            label = '表结构维护(变更)'
        elif has_select:
            label = '变量赋值/参数取值'
        else:
            label = '未识别'

    if has_grant:
        label += '+权限授予'
    return label


def extract_source_tables(content):
    """
    提取 FROM / JOIN 后面的 schema.table，保留原始大小写。
    注意：
      1. DELETE FROM <目标表> 这种写法里的表是被清空的目标表本身，不是来源表，
         因此这里显式排除 DELETE FROM 的情况，只保留 SELECT ... FROM 和 JOIN。
      2. 部分手写脚本里会直接写 $GDMDB.表名 这种不带花括号的变量格式，
         这种格式不会被 substitute_db_vars 统一替换（只处理 ${VAR} 格式），
         所以这里的正则允许 schema 前面带一个可选的 "$" 符号，
         从而能把 $GDMDB.表名 原样识别并输出，方便人工核查这类手写写法。
    """
    pattern = r'(DELETE\s+FROM|FROM|JOIN)\s+(\$?\w+\.\w+)'
    results = []
    for keyword, table in re.findall(pattern, content, flags=re.IGNORECASE):
        if keyword.strip().upper().startswith('DELETE'):
            continue
        results.append(table)
    return results


def filter_source_tables(table_names, english_name):
    """
    过滤来源表：
      - 去掉与目标表同名的自引用
      - 去掉 VT_{英文表名} 开头的临时表（这是脚本内的临时表，不作为真实来源表）
      - 按 (schema,table) 大小写不敏感去重，保留第一次出现的写法
    """
    exclude_prefix = f"VT_{english_name.upper()}"
    seen = set()
    filtered = []
    for table_name in table_names:
        schema, _, table_only = table_name.partition('.')
        table_only_upper = table_only.upper()
        if table_only_upper == english_name.upper():
            continue
        if table_only_upper.startswith(exclude_prefix):
            continue
        dedup_key = (schema.upper(), table_only_upper)
        if dedup_key in seen:
            continue
        seen.add(dedup_key)
        filtered.append((schema, table_only))
    return filtered


def process_file(file_path, file_name):
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        raw_content = f.read()

    comment, developed_by = extract_header_info(file_path)

    var_map = extract_db_var_map(raw_content)
    content = substitute_db_vars(raw_content, var_map)

    english_name, target_schema, target_type = determine_target_info(content, file_name)
    process_type = determine_process_type(content, target_type)

    source_tables_raw = extract_source_tables(content)
    source_tables = filter_source_tables(source_tables_raw, english_name)

    rows = []
    if source_tables:
        for source_schema, source_table in source_tables:
            rows.append((file_name, target_schema.upper(), english_name, comment, developed_by,
                         source_schema.upper(), source_table.upper(), process_type))
    else:
        # 没有解析到来源表也保留一条记录，方便人工核查该脚本
        rows.append((file_name, target_schema.upper(), english_name, comment, developed_by,
                     '', '', process_type))
    return rows


def process_folder(folder_path):
    data = []
    pl_files = []
    for root, dirs, files in os.walk(folder_path):
        for file_name in files:
            if file_name.lower().endswith(".pl"):
                pl_files.append((root, file_name))

    total_files = len(pl_files)
    for idx, (root, file_name) in enumerate(pl_files, start=1):
        file_path = os.path.join(root, file_name)
        try:
            data.extend(process_file(file_path, file_name))
        except Exception as e:
            print(f"\n处理{file_path}失败: {e}")
        progress_bar(idx, total_files, 30)
    print()
    return data


def write_to_excel(data, output_file):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    headers = ['脚本文件名', '目标库', '英文表名', '中文表名', '开发人员', '来源库名', '来源表名', '处理类型']
    for col_idx, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(data, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=value)

    widths = [32, 12, 30, 40, 12, 12, 40, 22]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    workbook.save(output_file)


def progress_bar(current, total, bar_length):
    progress = current / total if total else 1
    block = int(bar_length * progress)
    percentage = progress * 100
    text = (f"\r---------------------- 处理文件中 | "
            f"[{'#' * block}{'-' * (bar_length - block)}] |【{percentage:.2f}%】 ----------------------")
    sys.stdout.write(text)
    sys.stdout.flush()


def main(folder_path, output_file):
    data = process_folder(folder_path)
    print("---------------------- 将结果写入文件... ----------------------")
    write_to_excel(data, output_file)


if __name__ == "__main__":
    if len(sys.argv) == 2:
        folder_path = sys.argv[1]
        output_file = "table_rely_output_list.xlsx"
        main(folder_path, output_file)
        print(f"---------------------- 文件保存在:{output_file} ----------------------")
    else:
        print("用法: python get_rely_table_agl.py <scripts_folder_path>")