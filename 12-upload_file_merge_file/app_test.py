import os
import threading
import time
from flask import Flask, request, jsonify, send_file, render_template
from werkzeug.utils import secure_filename
import openpyxl

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])

processing = False  # 用于避免多个用户同时运行脚本
progress = 0  # 记录进度
lock = threading.Lock()

def log_error(log_file, message):
    """记录错误日志并打印"""
    with open(log_file, 'a', encoding='utf-8') as log:
        log.write(message + '\n')
    print(message)

def update_progress(value):
    """更新进度"""
    global progress
    with lock:
        progress = value

def reset_progress():
    """重置进度为0"""
    global progress
    with lock:
        progress = 0

def run_script(txt_file_path, xlsx_file_path):
    """运行脚本主逻辑"""
    global progress
    try:
        update_progress(10)  # 初始进度

        # 读取 txt 文件
        with open(txt_file_path, 'r', encoding='utf-8') as txt_file:
            table_names = [line.strip() for line in txt_file.readlines()]
        if not table_names:
            raise ValueError("TXT 文件为空")

        update_progress(30)  # 进度 30%

        # 打开 Excel 文件
        src_wb = openpyxl.load_workbook(xlsx_file_path)
        src_ws = src_wb.active

        # 创建新 Excel
        new_wb = openpyxl.Workbook()
        new_ws = new_wb.active

        # 筛选匹配行
        headers = [cell.value for cell in src_ws[1]]  # 假设第一行为表头
        new_ws.append(headers)
        for row in src_ws.iter_rows(values_only=True):
            if row[0] in table_names:  # 假设第一列为表名
                new_ws.append(row)

        update_progress(70)  # 进度 70%

        # 保存合并后的 Excel
        output_path = os.path.join(app.config['UPLOAD_FOLDER'], 'merged_pub_cd_map.xlsx')
        new_wb.save(output_path)
        src_wb.close()
        update_progress(100)  # 完成
        return output_path

    except Exception as e:
        log_error(os.path.join(app.config['UPLOAD_FOLDER'], 'error.log'), str(e))
        raise e

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_files():
    global processing
    if processing:
        return jsonify({"error": "脚本正在运行，请稍后再试"}), 429

    txt_file = request.files.get('txtFile')
    xlsx_file = request.files.get('xlsxFile')

    if not txt_file or not xlsx_file:
        return jsonify({"error": "请上传两个文件"}), 400

    # 保存上传的文件
    txt_file_path = os.path.join(app.config['UPLOAD_FOLDER'], secure_filename(txt_file.filename))
    xlsx_file_path = os.path.join(app.config['UPLOAD_FOLDER'], secure_filename(xlsx_file.filename))

    txt_file.save(txt_file_path)
    xlsx_file.save(xlsx_file_path)

    # 开启后台线程运行脚本
    def background_task():
        global processing
        try:
            run_script(txt_file_path, xlsx_file_path)
        finally:
            processing = False  # 任务完成后释放锁

    processing = True
    threading.Thread(target=background_task).start()

    return jsonify({"message": "文件已上传，正在处理"}), 202

@app.route('/progress', methods=['GET'])
def get_progress():
    return jsonify({"progress": progress})

@app.route('/download', methods=['GET'])
def download_file():
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], 'merged_pub_cd_map.xlsx')
    if not os.path.exists(file_path):
        return jsonify({"error": "文件尚未生成"}), 404
    return send_file(file_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
