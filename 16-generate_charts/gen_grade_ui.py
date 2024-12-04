import os
from tkinter import Tk, filedialog, StringVar, ttk, messagebox
from ttkthemes import ThemedTk
from openpyxl import load_workbook
from pyecharts.charts import Line
from pyecharts import options as opts
import pandas as pd
from datetime import datetime

# 主程序逻辑
#pyinstaller --noconfirm --onefile --windowed --add-data "themes;themes" --icon="tools.ico" trend_chart_tool.py

class TrendChartApp:
    def __init__(self, root):
        self.root = root
        self.root.title("成绩单趋势生成工具")
        self.root.geometry("600x400")
        self.root.configure(bg='#f0f0f0')  # 设置背景颜色
        self.root.set_theme("arc") #breeze
        self.root.option_add("*Font", "黑体 10")  # 设置全局字体
        
        # 初始化变量
        self.file_path = StringVar()
        self.sheet_name = StringVar()
        self.sheets = []

        # 文件选择框
        ttk.Label(root, text="选择要生成的工作簿:").pack(pady=10)
        file_frame = ttk.Frame(root)
        file_frame.pack(pady=5)
        ttk.Entry(file_frame, textvariable=self.file_path, width=40).pack(side="left", padx=5)
        ttk.Button(file_frame, text="选择文件", command=self.select_file).pack(side="left")

        # Sheet选择框
        ttk.Label(root, text="选择你要生成的工作表:").pack(pady=10)
        self.sheet_dropdown = ttk.Combobox(root, textvariable=self.sheet_name, state="readonly", width=30)
        self.sheet_dropdown.pack(pady=5)

        # 进度条
        #ttk.Label(root, text="生成进度:").pack(pady=10)
        self.progress = ttk.Progressbar(root, orient="horizontal", length=400, mode="determinate")
        self.progress.pack(pady=5)

        # 执行按钮
        ttk.Button(root, text="执行生成", command=self.generate_charts).pack(pady=20)

    def select_file(self):
        # 选择Excel文件
        file_path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx")])
        if file_path:
            self.file_path.set(file_path)
            self.load_sheets()

    def load_sheets(self):
        # 加载文件中的Sheet
        try:
            wb = load_workbook(self.file_path.get())
            self.sheets = wb.sheetnames
            self.sheet_dropdown["values"] = self.sheets
            if self.sheets:
                self.sheet_name.set(self.sheets[0])
        except Exception as e:
            messagebox.showerror("错误", f"无法加载文件: {e}")

    def generate_charts(self):
        # 检查输入
        if not self.file_path.get():
            messagebox.showwarning("警告", "请先选择文件！")
            return
        if not self.sheet_name.get():
            messagebox.showwarning("警告", "请先选择Sheet！")
            return
        
        # 执行折线图生成
        try:
            self.progress["value"] = 0
            self.root.update()
            self.run_script()
        except Exception as e:
            messagebox.showerror("错误", f"生成失败: {e}")

    def run_script(self):
        # 生成折线图的脚本
        excel_path = self.file_path.get()
        sheet_name = self.sheet_name.get()

        # 加载Excel
        wb = load_workbook(excel_path)
        ws = wb[sheet_name]

        # 获取数据
        data = [row for row in ws.iter_rows(values_only=True)]
        header = data[0]  # 第一行标题
        content = data[1:]  # 其余行内容
        df = pd.DataFrame(content, columns=header)

        # 配置保存目录
        output_folder = "趋势图"
        os.makedirs(output_folder, exist_ok=True)
        today = datetime.now().strftime("%Y%m%d")
        total_rows = len(df)
        self.progress["maximum"] = total_rows

        x_labels = df.columns[1:]  # 横坐标标签
        for index, row in df.iterrows():
            name = row["姓名"]
            scores = row[1:].tolist()

            # 最大值与最小值
            max_score = max(scores)
            min_score = min(scores)
            max_idx = scores.index(max_score)
            min_idx = scores.index(min_score)

            # 绘制折线图
            line = (
                Line()
                .add_xaxis(x_labels.tolist())
                .add_yaxis("分数", scores, is_smooth=True, label_opts=opts.LabelOpts(is_show=True))
                .set_global_opts(
                    title_opts=opts.TitleOpts(title=f"{name} 成绩趋势"),
                    xaxis_opts=opts.AxisOpts(name="单元"),
                    yaxis_opts=opts.AxisOpts(name="分数", max_=100, min_=0),
                    tooltip_opts=opts.TooltipOpts(is_show=True),
                )
                .set_series_opts(
                    markpoint_opts=opts.MarkPointOpts(
                        data=[
                            opts.MarkPointItem(name="最大值", coord=[x_labels[max_idx], max_score], value=max_score),
                            opts.MarkPointItem(name="最小值", coord=[x_labels[min_idx], min_score], value=min_score),
                        ]
                    ),
                    markline_opts=opts.MarkLineOpts(
                        data=[
                            opts.MarkLineItem(type_="max", name="最大值"),
                            opts.MarkLineItem(type_="min", name="最小值"),
                        ]
                    ),
                )
            )

            # 保存图表
            output_file = os.path.join(output_folder, f"{name}_成绩单_趋势_{today}.html")
            line.render(output_file)

            # 更新进度条
            self.progress["value"] += 1
            self.root.update()

        messagebox.showinfo("完成", f"所有图表已生成！图表保存在文件夹: {output_folder}")


if __name__ == "__main__":
    root = ThemedTk(theme=False)  # 使用ttkthemes美化
    app = TrendChartApp(root)
    root.mainloop()

