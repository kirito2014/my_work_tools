#获取Excel文件差异信息，用于比对设计文档差异信息

import os
import sys
import re,sys
import openpyxl
import pandas as pd
import numpy as np
import string
from datetime import datetime,date
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


def column_index_to_string(index):
    string_index = ''
    #print(index)
    while index >= 0:
        remainder = index % 26
        string_index = string.ascii_uppercase[remainder] + string_index
        index = index // 26 - 1
        #print(string_index)
    return string_index

'''
def find_added_and_deleted_rows(df_a,df_b):
    df_a.reset_index(drop=True)
    df_b.reset_index(drop=True)

    mergedA = df_b.merge(df_a, on=df_b.columns[1],how = 'left',indicator=True)
    mergedB = df_b.merge(df_a, on=df_b.columns[1],how = 'right',indicator=True)

    added_rows = mergedA[mergedA['_merge'] == 'left_only'].drop(columns=['_merge'])
    #print(added_rows)
    added_rows_indices = added_rows.index.tolist()
    print(len(added_rows_indices))

    deleted_rows = mergedB[mergedB['_merge'] == 'right_only'].drop(columns=['_merge'])
    deleted_rows_indices = deleted_rows.index.tolist()
    print(len(deleted_rows_indices))
    #print(deleted_rows)
    return added_rows_indices,deleted_rows_indices
'''
def find_added_and_deleted_rows(df_a,df_b):

    df_a['_original_index'] = df_a.index
    df_b['_original_index'] = df_b.index

    #mergedA = df_b.merge(df_a, on=df_b.columns[1],how = 'left',indicator=True)
    #mergedB = df_b.merge(df_a, on=df_b.columns[1],how = 'right',indicator=True)
    joined = df_a.merge(df_b,indicator=True,how='outer')

    added_rows = joined[joined['_merge'] == 'right_only'].drop(columns=['_merge'])
    #added_rows = mergedA[mergedA['_merge'] == 'left_only'].drop(columns=['_merge'])
    #print(added_rows)
    added_rows_indices = added_rows['_original_index'].tolist()
    #print(len(added_rows_indices))

    deleted_rows = joined[joined['_merge'] == 'left_only'].drop(columns=['_merge'])
    #deleted_rows = mergedB[mergedB['_merge'] == 'right_only'].drop(columns=['_merge'])
    deleted_rows_indices = deleted_rows['_original_index'].tolist()
    #print(len(deleted_rows_indices))
    #print(deleted_rows)
    return added_rows_indices,deleted_rows_indices


def compare_sheets(file_a,file_b,sheet_name):
    #读取sheet页
    df_a = pd.read_excel(file_a,sheet_name=sheet_name)
    df_b = pd.read_excel(file_b,sheet_name=sheet_name)

    df_a.columns = df_a.columns.str.strip()
    df_b.columns = df_b.columns.str.strip()

    #读取file_b的信息
    df_a_info = pd.read_excel(file_a,sheet_name="6.表设计")
    df_b_info = pd.read_excel(file_b,sheet_name="6.表设计")

    table_en_name_b = df_b_info.at[4,df_b_info.columns[2]]
    table_en_name_a = df_a_info.at[4,df_a_info.columns[2]]

    table_cn_name_b = df_b_info.at[5,df_b_info.columns[2]]
    table_cn_name_a = df_a_info.at[5,df_a_info.columns[2]]

    #对比文件名是否修改
    #获取6.表设计 表英文名 变动信息

    name_changes = []
    if table_en_name_a != table_en_name_b:
        name_changes.append((f"英文表名表更:{table_en_name_a}-->{table_en_name_b}"))
    if table_cn_name_a != table_cn_name_b:
        name_changes.append((f"中文表名表更:{table_cn_name_a}-->{table_cn_name_b}"))


    added_rows_indices,deleted_rows_indices = find_added_and_deleted_rows(df_a,df_b)

    row_changes = [(f"2--新增:{len(added_rows_indices)} 行, 删除:{len(deleted_rows_indices)} 行")]
    added_rows = df_b.iloc[added_rows_indices]
    #added_rows=df_b[~df_b['_original_index'].isin(df_a['_original_index'])]


    deleted_rows = df_a.iloc[deleted_rows_indices]
    #deleted_rows=df_a[~df_a['_original_index'].isin(df_b['_original_index'])]


    df_a.reset_index(drop=True,inplace=True)
    df_b.reset_index(drop=True,inplace=True)

    new_rows = [(f"新增行：第 {idx + 1 } 行:{row.tolist()}") for idx, row in added_rows.iterrows()]
    removed_rows = [(f"删除行：第 {idx + 1 } 行:{row.tolist()}") for idx, row in deleted_rows.iterrows()]
    
    df_a['_source'] = 'A'
    df_b['_source'] = 'B'


    merged_df = pd.concat([df_a, df_b])
    duplicates = merged_df.duplicated(subset=df_a.columns[:-2],keep=False)
    diff_df = merged_df[~duplicates]
    
    diff_in_a = diff_df[diff_df['_source'] == 'A'].drop(columns=['_source'])
    diff_in_b = diff_df[diff_df['_source'] == 'B'].drop(columns=['_source'])
    
    diff_idx_a = diff_in_a.index.tolist()
    diff_idx_b = diff_in_b.index.tolist()

    columns_a = diff_in_a.columns.tolist()
    columns_b = diff_in_b.columns.tolist()

    max_col_index = string.ascii_uppercase.index('Y')
    print(max_col_index)

    diff_list_a = []
    for idx in diff_idx_a:
        original_index = df_a.loc[idx,'_original_index']
        if original_index not in added_rows['_original_index'].values:
            for col in columns_a:
                col_index = df_a.columns.get_loc(col)
                if col_index >= max_col_index:
                    continue
                old_value = df_a.loc[idx, col]
                new_value = df_b.loc[idx, col] if original_index in df_b['_original_index'].values else np.nan
                if pd.isna(old_value) and pd.isna(new_value):
                    continue
                if old_value != new_value:
                    col_str = column_index_to_string(df_a.columns.get_loc(col))
                    diff_list_a.append((original_index + 1,col_str,old_value,new_value))

    diff_list_b = []
    for idx in diff_idx_b:
        original_index = df_b.loc[idx,'_original_index']
        if original_index not in deleted_rows['_original_index'].values:
            for col in columns_b:
                col_index = df_b.columns.get_loc(col)
                if col_index >= max_col_index:
                    continue
                old_value = df_a.loc[idx, col] if original_index in df_a['_original_index'].values    else np.nan
                new_value = df_b.loc[idx, col] 
                if pd.isna(old_value) and pd.isna(new_value):
                    continue
                if old_value != new_value:
                    col_str = column_index_to_string(df_b.columns.get_loc(col))
                    #print(df_b.columns.get_loc(col),col_str)
                    diff_list_b.append((original_index + 1,col_str,old_value,new_value))


    diff_result = {
        "differencesA": [(f"<第{idx + 1}行第{col_str}列>: {old_value}  -->  {new_value}") for (idx,col_str,old_value,new_value) in diff_list_a],
        "differencesB": [(f"<第{idx + 1}行第{col_str}列>: {old_value}  -->  {new_value}") for (idx,col_str,old_value,new_value) in diff_list_b],
        "add_rows": new_rows,
        "removed_rows":removed_rows,
        "row_changes":row_changes,
        "name_changes":name_changes if name_changes else ["无变更"],
    }

    return diff_result

#写入文件

def write_to_excel(file_a,file_b,sheet_name,diff_result):
    output_file = "聚合层设计文档对比变更记录.xlsx"
    today_str = date.today().strftime('%Y-%m-%d')

    #读取file_b的信息
    df_a_info = pd.read_excel(file_a,sheet_name=sheet_name)
    df_b_info = pd.read_excel(file_b,sheet_name=sheet_name)

    table_en_name_b = df_b_info.at[4,df_b_info.columns[2]]
    table_cn_name_b = df_b_info.at[5,df_b_info.columns[2]]

    #文件是否存在

    if os.path.exists(output_file):
        writer = pd.ExcelWriter(output_file,engine='openpyxl',mode='a')
        book = writer.book
        if "Sheet1" in book.sheetnames:
            sheet = book["Sheet1"]
        else:
            sheet = book.create_sheet("Sheet1")
    else:
        writer = pd.ExcelWriter(output_file,engine='openpyxl')
        book = writer.book
        sheet = book.create_sheet("Sheet1")

        sheet.append(["文件名","表英文名","表中文名","变更记录","新增行","删除行","其他变更","对比日期"])

    #删除已有的对比数据
    for row in sheet.iter_rows(min_row=2,max_row=sheet.max_row):
        if row[0].value == file_b and row[7].value == today_str:
            sheet.delete_rows(row[0].row,1)
    #写入数据
    change_logs = "\n".join(diff_result["differencesB"]) if diff_result["differencesB"] else "无变更"
    new_rows = "\n".join(diff_result["add_rows"]) if diff_result["add_rows"] else "无变更"
    removed_rows = "\n".join(diff_result["removed_rows"]) if diff_result["removed_rows"] else "无变更"
    other_changes = "\n".join(diff_result["name_changes"])
    

    sheet.append([file_b,table_en_name_b,table_cn_name_b,change_logs,new_rows,removed_rows,other_changes,today_str])

    for col in sheet.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) if (max_length + 2) < 25 else 25
        sheet.column_dimensions[col_letter].width = adjusted_width
        for cell in col:
            cell.alignment = Alignment(horizontal='left', vertical='center')

    book.save(output_file)
    writer.close()
    return other_changes


if __name__ == "__main__":
    #file_a = 'D:\\sunline_etl_tool\\get_excel_differences\\A.xlsx'
    #print(file_a)
    #file_b = 'D:\\sunline_etl_tool\\get_excel_differences\\B.xlsx'
    #print(file_b)
    #sheet_name = 'Sheet1'
     
    sheet_name = '7.字段设计'
    if len(sys.argv) != 3:
        print("Usage: python get_file_differences.py <original_file> <new_file>")
        sys.exit(1)
    file_a = sys.argv[1]
    file_b = sys.argv[2]
    output_file = "聚合层设计文档对比变更记录.xlsx"
    print (file_a)
    print (file_b) 

    diff_result = compare_sheets(file_a,file_b,sheet_name)
    write_to_excel(file_a,file_b,"6.表设计",diff_result)

    print("\n------------------ 以下是文件内容变更明细 -----------------\n")
    if diff_result["differencesB"]:
        print("1--内容变更：")
        for idx, item in enumerate(diff_result['differencesB'],1):
            print(f"    1.{idx}.{item}")

    print(diff_result['row_changes'][0])

    if diff_result["add_rows"]:
        print(" 2.1--新增字段变更：")
        for idx, item in enumerate(diff_result['add_rows'],1):
            print(f" {idx}、{item}")
    if diff_result["removed_rows"]:
        print(" 2.2--删除字段变更：")
        for idx, item in enumerate(diff_result['removed_rows'],1):
            print(f" {idx}、{item}")
    if diff_result["name_changes"]:
        print("3--表名变更：")
        for idx, item in enumerate(diff_result['name_changes'],1):
            print(f" 3.{idx}、{item}")
    print(f"\n------------------ 同步记录保存到文件{output_file}中 ------------------")