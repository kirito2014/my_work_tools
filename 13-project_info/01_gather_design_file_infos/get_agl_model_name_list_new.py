import os
import time
import sys
import  openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles import Font
from datetime import datetime
import xlrd
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd

print(openpyxl.__version__)
print(pd.__version__)

def get_data_from_excel(file_path):

    try:
        table_info_dict = []
        design_data = pd.read_excel(file_path, sheet_name = '6.表设计', usecols="A,C")
        design_data.columns = ['content','table_infos']
        design_data = design_data[1:]
        #print(design_data)
        for index,row in design_data.iterrows():
            row_dict = {}
            if index <= 9:
                content = row.iloc[0]
                table_infos = row.iloc[1]
            if pd.isna(table_infos):
                table_infos = '/'
            if isinstance(content,str) and content.startswith('主题对象'):
                row_dict['theme_object']=table_infos
            elif isinstance(content,str) and content.startswith('主题域分类（一级）'):
                row_dict['lvl_one_theme']=table_infos
            elif isinstance(content,str) and content.startswith('主题域分类（二级）'):
                row_dict['lvl_two_theme']=table_infos
            elif isinstance(content,str) and content.startswith('主题聚合英文名称'):
                row_dict['table_en_name']=table_infos
            elif isinstance(content,str) and content.startswith('主题聚合中文名称'):
                row_dict['table_ch_name']=table_infos
            elif isinstance(content,str) and content.startswith('主题聚合标签'):
                row_dict['table_tag']=table_infos
            elif isinstance(content,str) and content.startswith('主题聚合粒度'):
                row_dict['table_guanularity']=table_infos
            elif isinstance(content,str) and content.startswith('数据范围描述'):
                row_dict['theme_desc']=table_infos
            elif isinstance(content,str) and content.startswith('数据口径描述'):
                row_dict['theme_caliber']=table_infos
            table_info_dict.append(row_dict)
        return table_info_dict
    except Exception as e:
        print(f"[ ERROR ] 打开文件{os.path.basename(file_path)}失败:{e}.")

def merge_dicts(dict_list):
    merged_dict = {}
    for d in dict_list:
        merged_dict.update(d)
    return [merged_dict]

def extract_data_from_excel(file_path):

    if file_path.endswith(".xls") or file_path.endswith(".xlsx"):
        data_list = get_data_from_excel(file_path)
        #print(data_list)
        #data_dict = {}
        if data_list:
            data_dict = merge_dicts(data_list)
        data = data_dict[0]
        return data
    else:
        raise ValueError(f"Unsupported file type:{file_path}")

def extract_name_from_filename(file_name):
    base_name = os.path.splitext(file_name)[0] #去除拓展名
    last_dash_index = base_name.rfind('-')
    last_underscore_index = base_name.rfind('_')

    #取最大的索引值，即最右边的分隔符
    last_separator_index = max(last_dash_index,last_underscore_index)
    #提取名字部分
    if last_separator_index != -1:
        name = base_name[last_separator_index + 1:]
    else:
        name = base_name 
    return name 

def process_file(file_path):
    data = extract_data_from_excel(file_path)
    #print("正在处理:" + data.get("C7"))
    if data:
        file_name = os.path.basename(file_path)
        last_modified_time = datetime.fromtimestamp(os.path.getmtime(file_path)).strftime('%Y-%m-%d')
        #如果batch_number为“第”开头，则取batch_number为“第”开头的部分 否则为""
        # 解析文件名结构
        parts = file_name.split('-')
        batch_number = " "
        batch_group = " "

        if parts:
            # 检查第一个部分是否是批次号（以"第"开头且包含"批"）
            if parts[0].startswith('第') and '批' in parts[0]:
                batch_number = parts[0]
                # 组别为第二个部分（如果存在）
                batch_group = parts[1] if len(parts) >= 2 else " "
            else:
                # 没有批次号时，组别为第一个部分
                batch_group = parts[0]

        name = extract_name_from_filename(file_name)
        
        return (
            file_name,
            data.get("theme_object"), #主题对象
            data.get("lvl_one_theme"), #一级主题
            data.get("lvl_two_theme"), #二级主题
            data.get("table_en_name"), #表英文名
            data.get("table_ch_name"), #表中文名
            data.get("table_tag"), #主题聚合标签
            data.get("table_guanularity"), #主题聚合粒度
            data.get("theme_desc"), #数据范围描述
            data.get("theme_caliber"), #数据口径描述
            last_modified_time, #最后修改时间
            batch_number,        #批次号
            batch_group,         #归属组别
            name                 #设计人员名称
        )
        return None

def progress_bar(current,total,bar_length=50):
    progress=current / total
    block = int(bar_length * progress)
    percentage = progress * 100
    text = f"\r[ INFO ] | 处理文件中: [{'#' * block}{'-' * (bar_length - block)}] |【{percentage:.2f}%】"
    sys.stdout.write(text)
    sys.stdout.flush()

def main(folder_path):
    results = []
    file_paths = []

    for root, _, files in os.walk(folder_path):
        for file in files:
            #print(split(file,".")[0])
            if (file.endswith(".xls") or file.endswith(".xlsx")) and not file.startswith("~$"):
                file_path = os.path.join(root, file)
                file_paths.append(file_path)

    total_files = len(file_paths)
    processed_files = 0


    with ThreadPoolExecutor(max_workers=4) as executor:
        future_to_file = {executor.submit(process_file,file_path): file_path for file_path in file_paths}

        for future in as_completed(future_to_file):
            result = future.result()
            if result:
                results.append(result)
                
            processed_files += 1
            progress_bar(processed_files,total_files)
    print()


    #提取文件夹名称作为文件名的一部分
    folder_name = os.path.basename(os.path.normpath(folder_path))
    output_file_name = f"模型内容获取_{folder_name}.xlsx"

    #保存到新文件
    script_path = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(script_path,output_file_name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    headers = ["文件名","主题对象","一级主题","二级主题","英文表名","中文表名","主题聚合标签","主题聚合粒度","数据范围描述","数据口径描述","最新更新日期","归属批次","归属组别","设计人员"]
    header_fill = PatternFill(start_color='2C7B1F',end_color='2C7B1F',fill_type='solid')
    header_font = Font(color="FFFFFF",size=10,bold=True)
    ws.append(headers)

    for col_num,header in enumerate(headers,1):
        cell = ws.cell(row=1,column=col_num,value=header)
        cell.fill = header_fill
        cell.font = header_font
        ws.column_dimensions[cell.column_letter].width = 20

    for result in results:
        ws.append(result)

    wb.save(output_path)
    print(f"[ INFO ] 已获取文件数量:< {len(results)} > 个 ")
    print(f"[ INFO ] 文件保存路径:< {output_path} > ")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("usage: python get_agl_model_name_list.py <file_path>")
        sys.exit(1)
    folder_path = sys.argv[1]
    #folder_path = "D:\\sunline_etl_tool\\test\\"
    print(f"[ INFO ] 获取文件路径成功:< {folder_path} >")
    main(folder_path)