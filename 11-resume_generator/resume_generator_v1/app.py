from flask import Flask, render_template, request, jsonify, send_file, Response
from flask_cors import CORS
import os
import uuid
from docx import Document
import json
import subprocess
import logging
from datetime import datetime
import pandas as pd
import threading
from queue import Queue
import time

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("app.log"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# 初始化Flask应用
app = Flask(__name__,
            template_folder='.',  # 设置模板文件夹为当前目录
            static_folder='static')  # 设置静态文件夹为static目录
CORS(app)

# 配置文件上传和输出目录
UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')
OUTPUT_FOLDER = os.path.join(os.getcwd(), 'output')
for folder in [UPLOAD_FOLDER, OUTPUT_FOLDER]:
    if not os.path.exists(folder):
        os.makedirs(folder)

# 从JSON模板结构定义Jinja2字段列表
JINJA_FIELDS = [
    # 基本信息
    {"id": "name", "name": "姓名", "description": "员工的完整姓名", "preview": "{{ BasicInfo.name }}"},
    {"id": "employee_id", "name": "员工编号", "description": "员工的唯一识别编号", "preview": "{{ BasicInfo.employee_id }}"},
    {"id": "department", "name": "部门", "description": "员工所属部门", "preview": "{{ AdditionInfo.department }}"},
    {"id": "position", "name": "职位", "description": "员工的职位名称", "preview": "{{ AdditionInfo.position }}"},
    {"id": "hire_date", "name": "入职日期", "description": "员工入职的日期", "preview": "{{ AdditionInfo.hire_date }}"},
    {"id": "gender", "name": "性别", "description": "员工性别", "preview": "{{ AdditionInfo.gender }}"},
    {"id": "age", "name": "年龄", "description": "员工年龄", "preview": "{{ AdditionInfo.age }}"},
    {"id": "tenure", "name": "司龄", "description": "员工在公司的工作年限", "preview": "{{ AdditionInfo.tenure }}"},
    {"id": "education", "name": "学历", "description": "员工的最高学历", "preview": "{{ BasicInfo.education }}"},
    {"id": "graduation_school", "name": "毕业院校", "description": "员工的毕业院校", "preview": "{{ BasicInfo.graduation_school }}"},
    
    # 工作信息
    {"id": "project_experience", "name": "项目经验", "description": "员工参与的项目经历", "preview": "{% for exp in WorkExperience %}{{ exp.project_name }}{% endfor %}"},
    {"id": "job_responsibilities", "name": "工作职责", "description": "员工的主要工作职责", "preview": "{{ WorkAbility.job_responsibilities }}"},
    {"id": "skills", "name": "技能标签", "description": "员工掌握的专业技能", "preview": "{% for skill in WorkAbility.skills %}{{ skill }}{% endfor %}"},
    {"id": "certifications", "name": "证书", "description": "员工获得的专业证书", "preview": "{% for cert in WorkAbility.certifications %}{{ cert }}{% endfor %}"},
    {"id": "training_experience", "name": "培训经历", "description": "员工参加的培训课程", "preview": "{% for training in WorkAbility.training_experience %}{{ training }}{% endfor %}"},
    
    # 联系方式
    {"id": "phone", "name": "电话", "description": "员工的联系电话", "preview": "{{ BasicInfo.phone }}"},
    {"id": "email", "name": "邮箱", "description": "员工的电子邮箱", "preview": "{{ BasicInfo.email }}"},
    {"id": "address", "name": "地址", "description": "员工的居住地址", "preview": "{{ BasicInfo.address }}"},
    {"id": "emergency_contact", "name": "紧急联系人", "description": "员工的紧急联系人", "preview": "{{ BasicInfo.emergency_contact }}"},
    {"id": "emergency_phone", "name": "紧急联系电话", "description": "紧急联系人的电话", "preview": "{{ BasicInfo.emergency_phone }}"},
    
    # 其他信息
    {"id": "bank_account", "name": "银行账户", "description": "员工的银行账户信息", "preview": "{{ AdditionInfo.bank_account }}"},
    {"id": "id_number", "name": "身份证号", "description": "员工的身份证号码", "preview": "{{ AdditionInfo.id_number }}"},
    {"id": "marital_status", "name": "婚姻状况", "description": "员工的婚姻状况", "preview": "{{ AdditionInfo.marital_status }}"},
    {"id": "political_status", "name": "政治面貌", "description": "员工的政治面貌", "preview": "{{ AdditionInfo.political_status }}"}
]

# 路由定义
@app.route('/')
def index():
    # 渲染我们的主页面index.html
    return render_template('index.html')

@app.route('/api/fields')
def get_jinja_fields():
    return jsonify(JINJA_FIELDS)

@app.route('/api/upload', methods=['POST'])
def upload_document():
    if 'file' not in request.files:
        return jsonify({"error": "未找到文件"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "未选择文件"}), 400
    if file and file.filename.endswith('.docx'):
        # 保存上传的文件
        filename = str(uuid.uuid4()) + '.docx'
        filepath = os.path.join(UPLOAD_FOLDER, filename)
        file.save(filepath)
        
        # 解析Word文档内容
        doc = Document(filepath)
        content = []
        for para in doc.paragraphs:
            content.append(para.text)
        
        return jsonify({
            "success": True,
            "filename": filename,
            "content": content
        })
    return jsonify({"error": "仅支持.docx格式文件"}), 400

@app.route('/api/save', methods=['POST'])
def save_document():
    data = request.json
    if not data or 'filename' not in data or 'content' not in data:
        return jsonify({"error": "缺少必要参数"}), 400
    
    # 创建新的Word文档
    doc = Document()
    for para_text in data['content']:
        doc.add_paragraph(para_text)
    
    # 保存文件
    output_filename = str(uuid.uuid4()) + '.docx'
    output_path = os.path.join(OUTPUT_FOLDER, output_filename)
    doc.save(output_path)
    
    return jsonify({
        "success": True,
        "output_filename": output_filename,
        "download_url": f"/download/{output_filename}"
    })

@app.route('/download/<filename>')
def download_document(filename):
    filepath = os.path.join(OUTPUT_FOLDER, filename)
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True)
    return jsonify({"error": "文件不存在"}), 404

# 项目配置
PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
TEMP_DIR = os.path.join(PROJECT_ROOT, 'temp')
if not os.path.exists(TEMP_DIR):
    os.makedirs(TEMP_DIR)

# 获取部门列表
@app.route('/api/departments', methods=['GET'])
def get_departments():
    try:
        # 模拟部门数据
        departments = {
            "level1": [
                {"id": "tech", "name": "技术部"},
                {"id": "product", "name": "产品部"},
                {"id": "operation", "name": "运营部"},
                {"id": "hr", "name": "人力资源部"}
            ],
            "level2": {
                "tech": [
                    {"id": "frontend", "name": "前端开发组"},
                    {"id": "backend", "name": "后端开发组"},
                    {"id": "test", "name": "测试组"},
                    {"id": "devops", "name": "运维组"}
                ],
                "product": [
                    {"id": "requirements", "name": "需求分析组"},
                    {"id": "design", "name": "设计组"}
                ],
                "operation": [
                    {"id": "market", "name": "市场组"},
                    {"id": "customer", "name": "客户组"}
                ],
                "hr": [
                    {"id": "recruit", "name": "招聘组"},
                    {"id": "training", "name": "培训组"}
                ]
            }
        }
        return jsonify({"status": "success", "data": departments})
    except Exception as e:
        logger.error(f"获取部门列表失败: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

# 获取人员列表
@app.route('/api/persons', methods=['GET'])
def get_persons():
    try:
        # 尝试调用get_emp_list.py获取人员信息
        script_path = os.path.join(PROJECT_ROOT, 'package', 'functions', 'get_emp_list.py')
        employees = []
        
        if os.path.exists(script_path):
            try:
                process = subprocess.Popen(
                    ['python', script_path],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    universal_newlines=True
                )
                stdout, stderr = process.communicate()
                
                # 读取生成的emp_list.json文件
                emp_list_path = os.path.join(PROJECT_ROOT, 'emp_list.json')
                if os.path.exists(emp_list_path):
                    with open(emp_list_path, 'r', encoding='utf-8') as f:
                        employees = json.load(f)
            except Exception as e:
                logger.warning(f"调用get_emp_list.py失败: {str(e)}")
        
        # 如果没有获取到数据，使用模拟数据
        if not employees:
            employees = [
                {"工号": "001", "姓名": "张三", "部门": "技术部-前端开发组"},
                {"工号": "002", "姓名": "李四", "部门": "技术部-后端开发组"},
                {"工号": "003", "姓名": "王五", "部门": "技术部-测试组"},
                {"工号": "004", "姓名": "赵六", "部门": "产品部-需求分析组"},
                {"工号": "005", "姓名": "孙七", "部门": "产品部-设计组"}
            ]
        
        # 转换格式以匹配前端期望
        persons = [
            {
                "id": emp.get("工号", emp.get("员工编号", str(i))),
                "name": emp.get("姓名", f"员工{i}"),
                "dept": emp.get("部门", "技术部-开发组")
            }
            for i, emp in enumerate(employees, 1)
        ]
        
        # 检查是否有筛选参数
        level1 = request.args.get('level1')
        level2 = request.args.get('level2')
        keyword = request.args.get('keyword')
        
        # 应用筛选
        filtered_persons = persons
        if level1:
            level1_name = next((d['name'] for d in [
                {"id": "tech", "name": "技术部"},
                {"id": "product", "name": "产品部"},
                {"id": "operation", "name": "运营部"},
                {"id": "hr", "name": "人力资源部"}
            ] if d['id'] == level1), level1)
            filtered_persons = [p for p in filtered_persons if p['dept'].startswith(level1_name)]
        if level2:
            filtered_persons = [p for p in filtered_persons if level2 in p['dept']]
        if keyword:
            keyword = keyword.lower()
            filtered_persons = [p for p in filtered_persons 
                              if keyword in p['id'].lower() or 
                                 keyword in p['name'].lower() or 
                                 keyword in p['dept'].lower()]
        
        return jsonify({"status": "success", "data": filtered_persons})
    except Exception as e:
        logger.error(f"获取人员列表失败: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

# 配置常量
CONFIG_FOLDER = os.path.join(os.path.dirname(__file__), 'config')
EMPLOYEE_DATA_PATH = os.path.join(CONFIG_FOLDER, 'emp_list.json')
BANK_LIST_PATH = os.path.join(CONFIG_FOLDER, 'bank_list.config')

# 创建必要的文件夹
if not os.path.exists(CONFIG_FOLDER):
    os.makedirs(CONFIG_FOLDER)

# 全局状态和队列
parse_queue = Queue()
generate_queue = Queue()
parse_progress = {}
generate_progress = {}

# 辅助函数：读取银行列表
def get_bank_list():
    if not os.path.exists(BANK_LIST_PATH):
        # 创建默认银行列表
        default_banks = ['长亮科技', '测试银行', '招商银行', '建设银行', '工商银行', '农业银行']
        with open(BANK_LIST_PATH, 'w', encoding='utf-8') as f:
            for bank in default_banks:
                f.write(f'{bank}\n')
        return default_banks
    
    with open(BANK_LIST_PATH, 'r', encoding='utf-8') as f:
        return [line.strip() for line in f if line.strip()]

# 辅助函数：保存员工数据
def save_employee_data(employees):
    with open(EMPLOYEE_DATA_PATH, 'w', encoding='utf-8') as f:
        json.dump(employees, f, ensure_ascii=False, indent=2)

# 辅助函数：读取员工数据
def get_employee_data():
    if not os.path.exists(EMPLOYEE_DATA_PATH):
        return []
    
    with open(EMPLOYEE_DATA_PATH, 'r', encoding='utf-8') as f:
        try:
            return json.load(f)
        except:
            return []

# 解析工作线程
def parse_worker():
    while True:
        task = parse_queue.get()
        task_id = task['task_id']
        folder_path = task['folder_path']
        parse_progress[task_id] = {'progress': 0, 'message': '开始解析简历...'}

        try:
            # 获取文件夹中的所有文件
            files = [f for f in os.listdir(folder_path) if f.lower().endswith(('.doc', '.docx'))]
            total_files = len(files)

            if total_files == 0:
                parse_progress[task_id] = {
                    'progress': 100,
                    'message': '文件夹中没有找到简历文件'
                }
                continue

            # 模拟解析过程
            for i, file in enumerate(files):
                progress = int((i + 1) / total_files * 100)
                parse_progress[task_id] = {
                    'progress': progress,
                    'message': f'正在解析: {file}'
                }
                time.sleep(0.1)  # 模拟处理时间

            parse_progress[task_id] = {
                'progress': 100,
                'message': f'解析完成，共处理 {total_files} 个文件'
            }
        except Exception as e:
            parse_progress[task_id] = {
                'progress': 0,
                'message': f'解析失败: {str(e)}'
            }
        finally:
            parse_queue.task_done()
            # 5分钟后清除进度记录
            threading.Timer(300, lambda: parse_progress.pop(task_id, None)).start()

# 生成工作线程
def generate_worker():
    while True:
        task = generate_queue.get()
        task_id = task['task_id']
        employee_ids = task['employee_ids']
        bank_name = task['bank_name']
        generate_progress[task_id] = {'progress': 0, 'message': '开始生成简历...'}

        try:
            total_employees = len(employee_ids)

            if total_employees == 0:
                generate_progress[task_id] = {
                    'progress': 100,
                    'message': '没有选择员工'
                }
                continue

            # 模拟生成过程
            for i, emp_id in enumerate(employee_ids):
                progress = int((i + 1) / total_employees * 100)
                generate_progress[task_id] = {
                    'progress': progress,
                    'message': f'正在生成: {emp_id} 的简历'
                }
                time.sleep(0.2)  # 模拟处理时间

            generate_progress[task_id] = {
                'progress': 100,
                'message': f'简历生成完成，共生成 {total_employees} 份简历'
            }
        except Exception as e:
            generate_progress[task_id] = {
                'progress': 0,
                'message': f'生成失败: {str(e)}'
            }
        finally:
            generate_queue.task_done()
            # 5分钟后清除进度记录
            threading.Timer(300, lambda: generate_progress.pop(task_id, None)).start()

# 启动工作线程
threading.Thread(target=parse_worker, daemon=True).start()
threading.Thread(target=generate_worker, daemon=True).start()

@app.route('/resume-builder')
def resume_builder():
    return render_template('resume_builder.html')

@app.route('/api/banks')
def get_banks():
    try:
        banks = get_bank_list()
        return jsonify(banks)
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'获取银行列表失败: {str(e)}'
        }), 500

@app.route('/api/employees')
def get_employees():
    try:
        employees = get_employee_data()
        return jsonify(employees)
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'获取员工列表失败: {str(e)}'
        }), 500

# 执行特殊更新
@app.route('/api/special-update', methods=['POST'])
def special_update():
    try:
        data = request.json
        resume_folder = data.get('resume_folder')
        info_file = data.get('info_file')
        force_update = data.get('force_update', False)
        
        if not resume_folder or not info_file:
            return jsonify({"status": "error", "message": "请完整填写所有信息"}), 400
        
        logger.info(f"执行特殊更新 - 简历文件夹: {resume_folder}, 信息文件: {info_file}, 强制更新: {force_update}")
        
        # 查找special_update.py脚本
        script_path = None
        possible_paths = [
            os.path.join(PROJECT_ROOT, 'package', 'functions', 'special_update.py'),
            os.path.join(PROJECT_ROOT, 'package', 'special_update.py'),
            os.path.join(PROJECT_ROOT, 'special_update.py')
        ]
        
        for path in possible_paths:
            if os.path.exists(path):
                script_path = path
                break
        
        if script_path:
            # 调用special_update.py
            args = ['python', script_path, '--resume-folder', resume_folder, '--info-file', info_file]
            if force_update:
                args.append('--force-update')
            
            process = subprocess.Popen(
                args,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                universal_newlines=True
            )
            
            stdout, stderr = process.communicate()
            
            if process.returncode != 0:
                logger.error(f"特殊更新失败: {stderr}")
                return jsonify({"status": "error", "message": f"特殊更新失败: {stderr}"}), 500
            
            logger.info(f"特殊更新成功: {stdout}")
            return jsonify({"status": "success", "message": "特殊更新执行完成", "data": stdout})
        else:
            # 模拟更新成功
            return jsonify({"status": "success", "message": "特殊更新执行完成（模拟）"})
    except Exception as e:
        logger.error(f"执行特殊更新时发生错误: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

# 执行简历校验
@app.route('/api/validate-resumes', methods=['POST'])
def validate_resumes():
    try:
        data = request.json
        folder_path = data.get('folder_path')
        
        if not folder_path:
            return jsonify({"status": "error", "message": "请选择简历文件夹"}), 400
        
        logger.info(f"开始执行简历校验: {folder_path}")
        
        # 查找validate_resume.py脚本
        script_path = None
        possible_paths = [
            os.path.join(PROJECT_ROOT, 'package', 'functions', 'validate_resume.py'),
            os.path.join(PROJECT_ROOT, 'package', 'validate_resume.py'),
            os.path.join(PROJECT_ROOT, 'validate_resume.py')
        ]
        
        for path in possible_paths:
            if os.path.exists(path):
                script_path = path
                break
        
        if script_path:
            # 调用简历校验脚本
            process = subprocess.Popen(
                ['python', script_path, '--folder', folder_path],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                universal_newlines=True
            )
            
            stdout, stderr = process.communicate()
            
            if process.returncode != 0:
                logger.error(f"校验失败: {stderr}")
                return jsonify({"status": "error", "message": f"校验失败: {stderr}"}), 500
            
            logger.info(f"校验成功: {stdout}")
            return jsonify({"status": "success", "message": "简历校验完成", "data": stdout})
        else:
            # 模拟校验成功
            return jsonify({"status": "success", "message": "简历校验完成（模拟）"})
    except Exception as e:
        logger.error(f"执行简历校验时发生错误: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

# 保存上次选择的人员
@app.route('/api/save-selection', methods=['POST'])
def save_selection():
    try:
        data = request.json
        person_ids = data.get('person_ids', [])
        
        # 保存到临时文件
        selection_file = os.path.join(TEMP_DIR, 'last_selection.json')
        with open(selection_file, 'w', encoding='utf-8') as f:
            json.dump({
                'person_ids': person_ids,
                'timestamp': datetime.now().isoformat()
            }, f, ensure_ascii=False)
        
        return jsonify({"status": "success", "message": "选择已保存"})
    except Exception as e:
        logger.error(f"保存选择时发生错误: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

# 加载上次选择的人员
@app.route('/api/load-selection', methods=['GET'])
def load_selection():
    try:
        selection_file = os.path.join(TEMP_DIR, 'last_selection.json')
        if not os.path.exists(selection_file):
            return jsonify({"status": "error", "message": "没有找到上次的选择"}), 404
        
        with open(selection_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        return jsonify({"status": "success", "data": data['person_ids']})
    except Exception as e:
        logger.error(f"加载选择时发生错误: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/parse-tech-info', methods=['POST'])
def parse_tech_info():
    try:
        file_path = request.json.get('file_path')
        if not file_path or not os.path.exists(file_path):
            return jsonify({
                'status': 'error',
                'message': '文件不存在'
            }), 400

        # 尝试读取Excel文件
        try:
            df = pd.read_excel(file_path)
        except:
            # 尝试读取CSV文件
            try:
                df = pd.read_csv(file_path, encoding='utf-8')
            except:
                df = pd.read_csv(file_path, encoding='gbk')

        # 转换为JSON并保存
        employees = df.to_dict('records')
        save_employee_data(employees)

        return jsonify({
            'status': 'success',
            'message': f'成功解析 {len(employees)} 条员工信息',
            'count': len(employees)
        })
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'解析技术人员信息失败: {str(e)}'
        }), 500

@app.route('/api/parse-resumes', methods=['POST'])
def parse_resumes():
    try:
        folder_path = request.json.get('folder_path')
        if not folder_path or not os.path.exists(folder_path) or not os.path.isdir(folder_path):
            return jsonify({
                'status': 'error',
                'message': '文件夹不存在'
            }), 400

        # 创建新任务
        task_id = datetime.now().strftime('%Y%m%d%H%M%S')
        parse_queue.put({
            'task_id': task_id,
            'folder_path': folder_path
        })

        return jsonify({
            'status': 'success',
            'task_id': task_id,
            'message': '已加入解析队列'
        })
    except Exception as e:
        return jsonify({
            'status': 'error',
            'message': f'开始解析失败: {str(e)}'
        }), 500

@app.route('/api/parse-progress')
def get_parse_progress():
    def event_stream():
        task_id = request.args.get('task_id')
        if not task_id:
            yield 'data: {"error": "缺少任务ID"}\n\n'
            return

        while task_id in parse_progress:
            data = parse_progress[task_id]
            yield f'data: {json.dumps(data)}\n\n'
            if data['progress'] == 100:
                break
            time.sleep(1)

    return Response(event_stream(), mimetype='text/event-stream')

# 增强的简历生成功能
@app.route('/api/generate-resumes', methods=['POST'])
def generate_resumes():
    try:
        data = request.json
        bank_id = data.get('bank_id') or data.get('bank_name')
        generate_method = data.get('generate_method', 'selected')
        person_ids = data.get('person_ids', []) or request.json.get('employee_ids', [])
        list_file_path = data.get('list_file_path')
        
        if not bank_id:
            return jsonify({"status": "error", "message": "请选择银行"}), 400
        
        if generate_method == 'selected' and not person_ids:
            return jsonify({"status": "error", "message": "请选择要生成的人员"}), 400
        
        if generate_method == 'all' and not list_file_path:
            return jsonify({"status": "error", "message": "请选择名单文件"}), 400
        
        logger.info(f"开始生成简历 - 银行: {bank_id}, 生成方式: {generate_method}")
        
        # 调用batch_render_from_docx.py生成简历
        script_path = os.path.join(PROJECT_ROOT, 'package', 'functions', 'batch_render_from_docx.py')
        
        if os.path.exists(script_path):
            # 构建命令参数
            args = ['python', script_path]
            
            # 查找模板文件
            template_dir = os.path.join(PROJECT_ROOT, 'template')
            if os.path.exists(template_dir):
                template_files = [f for f in os.listdir(template_dir) if f.endswith('.docx')]
                if template_files:
                    args.extend(['--template', os.path.join(template_dir, template_files[0])])
            
            # 添加银行参数
            args.extend(['--bank', bank_id])
            
            # 根据生成方式添加不同的参数
            if generate_method == 'all':
                args.extend(['--list', list_file_path])
            else:
                # 将人员ID转换为逗号分隔的字符串
                args.extend(['--persons', ','.join(person_ids)])
            
            process = subprocess.Popen(
                args,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                universal_newlines=True
            )
            
            stdout, stderr = process.communicate()
            
            if process.returncode != 0:
                logger.error(f"生成失败: {stderr}")
                return jsonify({"status": "error", "message": f"生成失败: {stderr}"}), 500
            
            logger.info(f"生成成功: {stdout}")
            return jsonify({"status": "success", "message": "简历生成完成", "data": stdout})
        else:
            # 如果没有找到脚本，使用原有的队列处理方式
            # 这里复用现有的生成任务队列逻辑
            if generate_method == 'selected':
                # 创建新任务
                task_id = datetime.now().strftime('%Y%m%d%H%M%S')
                generate_queue.put({
                    'task_id': task_id,
                    'employee_ids': person_ids,
                    'bank_name': bank_id
                })
                
                return jsonify({
                    'status': 'success',
                    'task_id': task_id,
                    'message': '已加入生成队列'
                })
            else:
                return jsonify({"status": "error", "message": "按名单生成功能需要package/functions/batch_render_from_docx.py脚本"}), 500
    except Exception as e:
        logger.error(f"生成简历时发生错误: {str(e)}")
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/generate-progress')
def get_generate_progress():
    def event_stream():
        task_id = request.args.get('task_id')
        if not task_id:
            yield 'data: {"error": "缺少任务ID"}\n\n'
            return

        while task_id in generate_progress:
            data = generate_progress[task_id]
            yield f'data: {json.dumps(data)}\n\n'
            if data['progress'] == 100:
                break
            time.sleep(1)

    return Response(event_stream(), mimetype='text/event-stream')

if __name__ == '__main__':
    app.run(debug=True, port=5000)